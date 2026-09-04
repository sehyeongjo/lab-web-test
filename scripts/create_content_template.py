from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "outputs" / "eusun-han-lab" / "Eusun-Han-Lab-content-template.xlsx"

NAVY = "173B45"
TEAL = "3F7776"
PALE_TEAL = "EAF3F1"
CREAM = "F6F1E8"
WHITE = "FFFFFF"
TEXT = "27383D"
MUTED = "607176"
LINE = "CCD8D7"
HIDDEN_FILL = "F2F2F2"

THIN_BOTTOM = Border(bottom=Side(style="thin", color=LINE))


SHEETS = {
    "Settings": {
        "headers": ["key", "value"],
        "rows": [
            ["lab_name", "Eusun Han's Lab"],
            ["tagline", "Curious minds, careful research, meaningful impact."],
            [
                "introduction",
                "Eusun Han's Lab is a collaborative research group exploring intelligent systems and the ways they can better understand and support people. Replace this sample introduction in the Settings sheet.",
            ],
            ["hero_image_url", ""],
            ["affiliation", "Example University · Sample content"],
            ["contact_email", "lab@example.edu"],
            ["address", "Research Building, Example University"],
            ["footer_text", "Eusun Han's Lab"],
            ["research_heading", "Research"],
            [
                "research_intro",
                "Our sample research program connects human-centered questions with reliable computational methods. Replace these topics with the lab's actual research areas.",
            ],
            [
                "publications_intro",
                "Selected sample publications are listed below. Add the lab's work to the Publications sheet.",
            ],
        ],
        "widths": [27, 88],
    },
    "News": {
        "headers": ["month", "content", "link_label", "link_url", "visible"],
        "rows": [
            [
                datetime(2026, 9, 1),
                "Sample news: the new lab website is ready for content.",
                "",
                "",
                True,
            ],
            [
                datetime(2026, 6, 1),
                "Sample news: add a new member announcement in the News sheet.",
                "",
                "",
                True,
            ],
            [
                datetime(2026, 3, 1),
                "Sample news: link a paper, award, or event from this row.",
                "Example link",
                "https://example.com",
                True,
            ],
            [
                datetime(2026, 1, 1),
                "This example row is hidden from the website because visible is FALSE.",
                "",
                "",
                False,
            ],
        ],
        "widths": [16, 72, 20, 42, 12],
    },
    "Professor": {
        "headers": [
            "name",
            "secondary_name",
            "photo_url",
            "title",
            "affiliation",
            "email",
            "personal_url",
            "cv_url",
            "scholar_url",
            "bio",
            "visible",
        ],
        "rows": [
            [
                "Eusun Han",
                "Professor profile · Sample content",
                "",
                "Professor",
                "Department Name, Example University",
                "eusun.han@example.edu",
                "https://example.com",
                "https://example.com",
                "https://scholar.google.com",
                "This is sample profile text. Add the professor's biography, interests, and current appointments in the Professor sheet.",
                True,
            ]
        ],
        "widths": [24, 30, 42, 22, 36, 28, 42, 42, 42, 82, 12],
    },
    "ProfessorDetails": {
        "headers": [
            "section",
            "display_order",
            "content",
            "subtext",
            "link_label",
            "link_url",
            "visible",
        ],
        "rows": [
            [
                "Education",
                1,
                "Ph.D. in Your Field, Example University",
                "2014–2019 · Sample content",
                "",
                "",
                True,
            ],
            [
                "Education",
                2,
                "B.S. in Your Field, Example University",
                "2010–2014 · Sample content",
                "",
                "",
                True,
            ],
            [
                "Selected Publications",
                1,
                "Example Paper: A Clear Title for a Representative Publication",
                "Sample Conference, 2026",
                "Paper",
                "https://example.com",
                True,
            ],
            ["Reviewer", 1, "Example conferences and journals", "Sample content", "", "", True],
            ["Awards", 1, "Example Research Award", "2025 · Sample content", "", "", True],
            ["Talks", 1, "Example invited talk", "Sep 2026 · Sample content", "", "", True],
        ],
        "widths": [26, 15, 68, 38, 18, 42, 12],
    },
    "Students": {
        "headers": [
            "group",
            "display_order",
            "name",
            "secondary_name",
            "photo_url",
            "status",
            "affiliation",
            "email",
            "personal_url",
            "cv_url",
            "scholar_url",
            "research_topics",
            "visible",
        ],
        "rows": [
            [
                "Graduate",
                1,
                "Graduate Student 01",
                "Sample member",
                "",
                "M.S. Student",
                "Example University",
                "student01@example.edu",
                "",
                "",
                "",
                "Human-Centered AI, Machine Learning",
                True,
            ],
            [
                "Graduate",
                2,
                "Graduate Student 02",
                "Sample member",
                "",
                "Ph.D. Student",
                "Example University",
                "student02@example.edu",
                "",
                "",
                "",
                "Multimodal Learning",
                True,
            ],
            [
                "Undergraduate",
                1,
                "Undergraduate Student 01",
                "Sample member",
                "",
                "Undergraduate Researcher",
                "Example University",
                "student03@example.edu",
                "",
                "",
                "",
                "Data-Efficient Learning",
                True,
            ],
        ],
        "widths": [18, 15, 26, 22, 42, 25, 34, 28, 42, 42, 42, 45, 12],
    },
    "Alumni": {
        "headers": [
            "display_order",
            "name",
            "secondary_name",
            "photo_url",
            "degree",
            "period",
            "current_position",
            "personal_url",
            "visible",
        ],
        "rows": [
            [
                1,
                "Alumni Example 01",
                "Sample member",
                "",
                "M.S.",
                "2023–2025",
                "Example Company",
                "",
                True,
            ]
        ],
        "widths": [15, 26, 22, 42, 14, 18, 38, 42, 12],
    },
    "Research": {
        "headers": [
            "display_order",
            "title",
            "summary",
            "details",
            "image_url",
            "image_alt",
            "link_label",
            "link_url",
            "visible",
        ],
        "rows": [
            [
                1,
                "Human-Centered AI",
                "Designing intelligent systems around real human needs.",
                "We study how people understand, use, and collaborate with intelligent systems.\nThis is sample content and should be replaced in the Research sheet.",
                "",
                "Sample placeholder for human-centered AI research",
                "",
                "",
                True,
            ],
            [
                2,
                "Data-Efficient Learning",
                "Learning useful representations from limited supervision.",
                "We explore methods that reduce the cost of data collection and annotation.\nAdd projects, methods, and applications relevant to the lab.",
                "",
                "Sample placeholder for data-efficient learning research",
                "",
                "",
                True,
            ],
            [
                3,
                "Multimodal Intelligence",
                "Connecting language, vision, and structured information.",
                "We investigate models that reason across different forms of information.\nReplace this sample area with the lab's actual research direction.",
                "",
                "Sample placeholder for multimodal intelligence research",
                "",
                "",
                True,
            ],
        ],
        "widths": [15, 30, 58, 78, 42, 52, 18, 42, 12],
    },
    "Publications": {
        "headers": [
            "category",
            "year",
            "display_order",
            "venue",
            "title",
            "authors",
            "paper_url",
            "project_url",
            "code_url",
            "video_url",
            "visible",
        ],
        "rows": [
            [
                "International Conference",
                2026,
                1,
                "Sample Conference 2026",
                "Example Paper: Replace This with a Publication Title",
                "Eusun Han and Sample Collaborators",
                "https://example.com",
                "",
                "https://example.com",
                "",
                True,
            ],
            [
                "Journal",
                2025,
                1,
                "Sample Journal",
                "Example Journal Article for the Google Sheet Template",
                "Sample Author, Eusun Han",
                "https://example.com",
                "",
                "",
                "",
                True,
            ],
        ],
        "widths": [28, 12, 15, 34, 72, 46, 42, 42, 42, 42, 12],
    },
}


HEADER_NOTES = {
    "group": "Type any group name. Website sections follow the order in which groups first appear in this sheet.",
    "visible": "FALSE hides this row from the website. Blank or TRUE displays it.",
    "display_order": "Smaller numbers appear first within the relevant group or section.",
    "month": "Enter a real date. The website displays it as Mon YYYY and sorts newest first.",
    "photo_url": "Public http:// or https:// image URL. Leave blank to use an initials placeholder.",
    "image_url": "Public http:// or https:// image URL. Leave blank to use a placeholder.",
    "link_url": "Use only http://, https://, or mailto: URLs.",
    "personal_url": "Use only http://, https://, or mailto: URLs.",
    "cv_url": "Use only http://, https://, or mailto: URLs.",
    "scholar_url": "Use only http://, https://, or mailto: URLs.",
    "paper_url": "Use only http://, https://, or mailto: URLs.",
    "project_url": "Use only http://, https://, or mailto: URLs.",
    "code_url": "Use only http://, https://, or mailto: URLs.",
    "video_url": "Use only http://, https://, or mailto: URLs.",
}


def add_list_validation(ws, cell_range: str, choices: list[str], prompt: str) -> None:
    validation = DataValidation(
        type="list",
        formula1='"' + ",".join(choices) + '"',
        allow_blank=True,
    )
    validation.promptTitle = "Choose a value"
    validation.prompt = prompt
    validation.errorTitle = "Invalid value"
    validation.error = "Select a value from the dropdown list."
    validation.errorStyle = "stop"
    validation.showErrorMessage = True
    validation.showInputMessage = True
    ws.add_data_validation(validation)
    validation.add(cell_range)


def add_order_validation(ws, cell_range: str) -> None:
    validation = DataValidation(
        type="whole",
        operator="between",
        formula1="1",
        formula2="999",
        allow_blank=True,
    )
    validation.promptTitle = "Display order"
    validation.prompt = "Enter a whole number from 1 to 999. Smaller numbers appear first."
    validation.error = "Enter a whole number from 1 to 999."
    validation.showErrorMessage = True
    validation.showInputMessage = True
    ws.add_data_validation(validation)
    validation.add(cell_range)


def add_year_validation(ws, cell_range: str) -> None:
    validation = DataValidation(
        type="whole",
        operator="between",
        formula1="1900",
        formula2="2100",
        allow_blank=True,
    )
    validation.error = "Enter a four-digit year from 1900 to 2100."
    validation.showErrorMessage = True
    ws.add_data_validation(validation)
    validation.add(cell_range)


def style_data_sheet(ws, headers: list[str], rows: list[list[object]], widths: list[int]) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(headers)).column_letter}{len(rows) + 1}"
    ws.row_dimensions[1].height = 32

    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos", size=11, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        note = HEADER_NOTES.get(str(cell.value))
        if note:
            cell.comment = Comment(note, "User")

    for col_index, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, col_index).column_letter].width = width

    for row in ws.iter_rows(min_row=2, max_row=len(rows) + 1, max_col=len(headers)):
        ws.row_dimensions[row[0].row].height = 46
        for cell in row:
            cell.font = Font(name="Aptos", size=10, color=TEXT)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = THIN_BOTTOM

    for header_index, header in enumerate(headers, start=1):
        column_letter = ws.cell(1, header_index).column_letter
        if header == "month":
            for cell in ws[column_letter][1:]:
                cell.number_format = "mmm yyyy"
        elif header in {"display_order", "year"}:
            for cell in ws[column_letter][1:]:
                cell.alignment = Alignment(horizontal="center", vertical="top")
        elif header == "visible":
            for cell in ws[column_letter][1:]:
                cell.alignment = Alignment(horizontal="center", vertical="top")


def add_sheet_validations(ws, headers: list[str]) -> None:
    positions = {header: ws.cell(1, index).column_letter for index, header in enumerate(headers, 1)}

    if "visible" in positions:
        col = positions["visible"]
        add_list_validation(ws, f"{col}2:{col}500", ["TRUE", "FALSE"], "TRUE or blank displays the row; FALSE hides it.")
        ws.conditional_formatting.add(
            f"A2:{col}500",
            FormulaRule(
                formula=[f"=${col}2=FALSE"],
                fill=PatternFill("solid", fgColor=HIDDEN_FILL),
                font=Font(color="8A8A8A", italic=True),
            ),
        )

    if "display_order" in positions:
        col = positions["display_order"]
        add_order_validation(ws, f"{col}2:{col}500")

    if "year" in positions:
        col = positions["year"]
        add_year_validation(ws, f"{col}2:{col}500")

    if ws.title == "ProfessorDetails":
        add_list_validation(
            ws,
            "A2:A500",
            ["Education", "Selected Publications", "Reviewer", "Awards", "Talks"],
            "Choose the professor profile section.",
        )
    elif ws.title == "Students":
        add_list_validation(
            ws,
            "F2:F500",
            ["Ph.D. Student", "M.S. Student", "Undergraduate Researcher"],
            "Choose the closest status or type a custom value after clearing validation.",
        )
    elif ws.title == "Publications":
        add_list_validation(
            ws,
            "A2:A500",
            ["International Conference", "Journal", "Workshop", "Preprint", "Other"],
            "Choose a publication category.",
        )


def build_instructions(ws) -> None:
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:F2")
    title = ws["A1"]
    title.value = "Eusun Han's Lab · Content Workbook"
    title.fill = PatternFill("solid", fgColor=NAVY)
    title.font = Font(name="Georgia", size=22, bold=True, color=WHITE)
    title.alignment = Alignment(vertical="center")
    for row in ws["A1:F2"]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=NAVY)
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 22

    ws.merge_cells("A4:F4")
    ws["A4"] = "QUICK START"
    ws["A4"].fill = PatternFill("solid", fgColor=TEAL)
    ws["A4"].font = Font(name="Aptos", size=11, bold=True, color=WHITE)

    steps = [
        ("1", "Upload this XLSX file to Google Drive and open it with Google Sheets."),
        ("2", "In Google Sheets, choose File → Share → Publish to web, then publish the entire document."),
        ("3", "Keep edit access restricted to lab administrators; public read access is sufficient for the website."),
        ("4", "Copy the Sheet ID from the URL: the text between /d/ and /edit."),
        ("5", "Set GOOGLE_SHEET_ID in .env.local, restart the local site, and refresh the browser."),
    ]
    ws["A5"] = "Step"
    ws["B5"] = "How to use"
    ws.merge_cells("B5:F5")
    for cell in ws["A5:F5"][0]:
        cell.fill = PatternFill("solid", fgColor=CREAM)
        cell.font = Font(name="Aptos", size=10, bold=True, color=TEXT)
        cell.border = THIN_BOTTOM
    for index, (number, instruction) in enumerate(steps, start=6):
        ws.cell(index, 1, number)
        ws.cell(index, 2, instruction)
        ws.merge_cells(start_row=index, start_column=2, end_row=index, end_column=6)
        ws.row_dimensions[index].height = 34
        for cell in ws[index]:
            cell.font = Font(name="Aptos", size=10, color=TEXT)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = THIN_BOTTOM
        ws.cell(index, 1).font = Font(name="Aptos", size=11, bold=True, color=TEAL)
        ws.cell(index, 1).alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A13:F13")
    ws["A13"] = "CONTENT RULES"
    ws["A13"].fill = PatternFill("solid", fgColor=TEAL)
    ws["A13"].font = Font(name="Aptos", size=11, bold=True, color=WHITE)

    rules = [
        ("Tab and header names", "Do not rename tabs or row-one headers. The website uses them as a fixed interface."),
        ("Visibility", "Only visible=FALSE hides a row. Blank values and TRUE values are displayed."),
        ("News dates", "Enter a real date in month. It is displayed as Sep 2026 and sorted newest first."),
        ("URLs", "Use only public http:// or https:// URLs. Link fields may also use mailto:. Invalid URLs are ignored."),
        ("Images", "Use a direct, publicly accessible image URL. A name-initial placeholder appears when blank or unavailable."),
        ("Long text", "Use plain text and line breaks. HTML and Markdown are intentionally not rendered."),
        ("Sample content", "Every populated row is an English example and should be replaced with the lab's verified information."),
    ]
    ws["A14"] = "Rule"
    ws["B14"] = "Details"
    ws.merge_cells("B14:F14")
    for cell in ws["A14:F14"][0]:
        cell.fill = PatternFill("solid", fgColor=CREAM)
        cell.font = Font(name="Aptos", size=10, bold=True, color=TEXT)
        cell.border = THIN_BOTTOM
    for index, (rule, details) in enumerate(rules, start=15):
        ws.cell(index, 1, rule)
        ws.cell(index, 2, details)
        ws.merge_cells(start_row=index, start_column=2, end_row=index, end_column=6)
        ws.row_dimensions[index].height = 38
        for cell in ws[index]:
            cell.font = Font(name="Aptos", size=10, color=TEXT)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = THIN_BOTTOM
        ws.cell(index, 1).font = Font(name="Aptos", size=10, bold=True, color=NAVY)

    ws.column_dimensions["A"].width = 24
    for column in "BCDEF":
        ws.column_dimensions[column].width = 18
    ws.freeze_panes = "A5"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1


def create_workbook() -> Path:
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "Instructions"
    build_instructions(instructions)

    for name, config in SHEETS.items():
        ws = workbook.create_sheet(name)
        headers = config["headers"]
        rows = config["rows"]
        ws.append(headers)
        for row in rows:
            ws.append(row)
        style_data_sheet(ws, headers, rows, config["widths"])
        add_sheet_validations(ws, headers)

    workbook.active = 0
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)
    return OUTPUT


def verify_workbook(path: Path) -> None:
    workbook = load_workbook(path, data_only=False)
    expected_names = ["Instructions", *SHEETS.keys()]
    assert workbook.sheetnames == expected_names

    for name, config in SHEETS.items():
        ws = workbook[name]
        headers = [cell.value for cell in ws[1]]
        assert headers == config["headers"], f"Header mismatch in {name}"
        assert ws.freeze_panes == "A2"
        assert ws.auto_filter.ref
        assert ws.max_row == len(config["rows"]) + 1
        assert len(ws.data_validations.dataValidation) > 0 or name in {"Settings", "Professor"}

    assert isinstance(workbook["News"]["A2"].value, datetime)
    assert workbook["News"]["A2"].number_format == "mmm yyyy"
    assert workbook["News"]["E5"].value is False
    assert workbook["Settings"]["A2"].value == "lab_name"
    assert workbook["Settings"]["B2"].value == "Eusun Han's Lab"


if __name__ == "__main__":
    output_path = create_workbook()
    verify_workbook(output_path)
    print(output_path)
